"""Enterprise source loaders for structured, semi-structured, and document inputs."""

from __future__ import annotations

import json
import re
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
from collections.abc import Mapping
from email import policy
from email.message import EmailMessage, Message
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from kg_construction.domain.models import SourceRecord
from kg_construction.infrastructure.loaders.common import (
    AdapterNotConfigured,
    attach_source_evidence,
    file_location,
    jsonable_value,
    normalize_rows,
)

JsonDict = dict[str, Any]

_CHAT_WITH_TIMESTAMP_RE = re.compile(
    r"^\[(?P<timestamp>[^\]]+)\]\s*(?P<speaker>[^:]+):\s*(?P<message>.*)$"
)
_CHAT_SPEAKER_RE = re.compile(r"^(?P<speaker>[A-Za-z][^:]{0,80}):\s*(?P<message>.*)$")
_MARKDOWN_HEADING_RE = re.compile(r"^(?P<level>#{1,6})\s+(?P<title>.+?)\s*$")
_PDF_TEXT_RE = re.compile(r"\((?P<text>(?:\\.|[^\\()])*)\)\s*Tj")
_HTML_SPACE_RE = re.compile(r"\s+")
_APACHE_LOG_RE = re.compile(
    r"^(?P<client_ip>\S+)\s+(?P<ident>\S+)\s+(?P<user>\S+)\s+"
    r"\[(?P<timestamp>[^\]]+)\]\s+"
    r"\"(?P<method>\S+)\s+(?P<url>\S+)(?:\s+(?P<protocol>[^\"]+))?\"\s+"
    r"(?P<status>\d{3})\s+(?P<bytes>\S+)"
    r"(?:\s+\"(?P<referrer>[^\"]*)\"\s+\"(?P<user_agent>[^\"]*)\")?$"
)
_KEY_VALUE_RE = re.compile(r"(?P<key>[A-Za-z_][A-Za-z0-9_.-]*)=(?P<value>\"[^\"]*\"|\S+)")
_GENERIC_LOG_RE = re.compile(
    r"^(?:(?P<timestamp>\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}:\d{2}(?:[,.]\d+)?Z?)\s+)?"
    r"(?:(?P<level>TRACE|DEBUG|INFO|WARN|WARNING|ERROR|CRITICAL)\s+)?"
    r"(?:(?P<logger>[A-Za-z0-9_.:-]+)\s+)?"
    r"(?P<message>.*)$",
    flags=re.IGNORECASE,
)
_DOI_RE = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b")
_TEMPLATE_NUMBER_RE = re.compile(r"\b\d+(?:\.\d+)?\b")
_TEMPLATE_HEX_RE = re.compile(r"\b0x[0-9A-Fa-f]+\b")
_TEMPLATE_UUID_RE = re.compile(r"\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b")
_TEMPLATE_ID_RE = re.compile(r"\b[A-Z]+-\d+\b")


def _optional_dependency(package: str, install_hint: str) -> AdapterNotConfigured:
    return AdapterNotConfigured(f"{package} is required for this loader. {install_hint}")


def _load_json_payload(path: Path) -> Any:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def _source_name(path: str | Path, fragment: str | None = None) -> str:
    name = str(Path(path))
    return f"{name}#{fragment}" if fragment else name


def _collapse_space(value: str) -> str:
    return _HTML_SPACE_RE.sub(" ", value).strip()


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value not in (None, "") else f"column_{index}"
        base = re.sub(r"\s+", "_", base).strip("_") or f"column_{index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _is_empty_row(values: tuple[Any, ...]) -> bool:
    return all(value in (None, "") for value in values)


class ExcelLoader:
    """Load each Excel worksheet as a structured source record."""

    name = "excel_loader"
    source_type = "excel"

    def load(self, path: str | Path) -> list[SourceRecord]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - exercised when dependency is absent
            raise _optional_dependency("openpyxl", "Install kg-construction with Excel support.") from exc

        workbook_path = Path(path)
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        sources: list[SourceRecord] = []
        try:
            for worksheet in workbook.worksheets:
                headers: list[str] | None = None
                header_row_number: int | None = None
                rows: list[JsonDict] = []
                row_numbers: list[int] = []

                for row_number, values in enumerate(
                    worksheet.iter_rows(values_only=True),
                    start=1,
                ):
                    if _is_empty_row(values):
                        continue
                    if headers is None:
                        headers = _unique_headers(list(values))
                        header_row_number = row_number
                        continue
                    row = {
                        header: jsonable_value(value)
                        for header, value in zip(headers, values, strict=False)
                    }
                    if not any(value not in (None, "") for value in row.values()):
                        continue
                    rows.append(row)
                    row_numbers.append(row_number)

                if headers is None:
                    continue

                source = SourceRecord.from_rows(
                    normalize_rows(rows),
                    source_type=self.source_type,
                    source_name=_source_name(workbook_path, worksheet.title),
                    metadata={
                        "loader": self.name,
                        "workbook_path": str(workbook_path),
                        "sheet_name": worksheet.title,
                        "header_row_number": header_row_number,
                        "row_numbers": row_numbers,
                        "row_count": len(rows),
                        "columns": headers,
                    },
                )
                attach_source_evidence(
                    source,
                    location=file_location(
                        workbook_path,
                        kind="excel_worksheet",
                        sheet_name=worksheet.title,
                        header_row_number=header_row_number,
                        row_count=len(rows),
                    ),
                    text=f"{worksheet.title}: {len(rows)} data rows",
                )
                sources.append(source)
        finally:
            workbook.close()

        return sources


class SQLDatabaseLoader:
    """Load relational tables or explicit SQL query results through SQLAlchemy."""

    name = "sql_database_loader"
    source_type = "sql"

    def __init__(
        self,
        *,
        connection_url: str | None = None,
        table: str | None = None,
        query: str | None = None,
        limit: int | None = None,
    ) -> None:
        self.connection_url = connection_url
        self.table = table
        self.query = query
        self.limit = limit

    def load(self, path: str | Path) -> list[SourceRecord]:
        try:
            from sqlalchemy import MetaData, Table, create_engine, inspect, select, text
        except ImportError as exc:  # pragma: no cover - exercised when dependency is absent
            raise _optional_dependency("SQLAlchemy", "Install kg-construction with SQL support.") from exc

        database_url = self.connection_url or f"sqlite:///{Path(path)}"
        engine = create_engine(database_url)
        sources: list[SourceRecord] = []
        try:
            inspector = inspect(engine)
            with engine.connect() as connection:
                if self.query:
                    query_statement = text(self.query)
                    if self.limit and "limit" not in self.query.casefold():
                        query_statement = text(f"{self.query.rstrip(';')} LIMIT {self.limit}")
                    rows = [dict(row) for row in connection.execute(query_statement).mappings()]
                    source = self._build_source(
                        normalize_rows(rows),
                        source_name=f"{database_url}#query",
                        metadata={
                            "loader": self.name,
                            "database_url": database_url,
                            "query": self.query,
                            "row_count": len(rows),
                        },
                        location={"kind": "sql_query", "database_url": database_url},
                    )
                    sources.append(source)
                    return sources

                table_names = [self.table] if self.table else inspector.get_table_names()
                metadata = MetaData()
                for table_name in table_names:
                    if not table_name:
                        continue
                    table_obj = Table(table_name, metadata, autoload_with=engine)
                    table_statement = select(table_obj)
                    if self.limit:
                        table_statement = table_statement.limit(self.limit)
                    rows = [dict(row) for row in connection.execute(table_statement).mappings()]
                    columns = [column["name"] for column in inspector.get_columns(table_name)]
                    primary_key = inspector.get_pk_constraint(table_name).get(
                        "constrained_columns",
                        [],
                    )
                    foreign_keys = [
                        {
                            "constrained_columns": item.get("constrained_columns", []),
                            "referred_table": item.get("referred_table"),
                            "referred_columns": item.get("referred_columns", []),
                        }
                        for item in inspector.get_foreign_keys(table_name)
                    ]
                    source = self._build_source(
                        normalize_rows(rows),
                        source_name=f"{database_url}#{table_name}",
                        metadata={
                            "loader": self.name,
                            "database_url": database_url,
                            "table": table_name,
                            "columns": columns,
                            "primary_key": primary_key,
                            "foreign_keys": foreign_keys,
                            "row_count": len(rows),
                        },
                        location={
                            "kind": "sql_table",
                            "database_url": database_url,
                            "table": table_name,
                        },
                    )
                    sources.append(source)
        finally:
            engine.dispose()
        return sources

    def _build_source(
        self,
        rows: list[JsonDict],
        *,
        source_name: str,
        metadata: JsonDict,
        location: JsonDict,
    ) -> SourceRecord:
        source = SourceRecord.from_rows(
            rows,
            source_type=self.source_type,
            source_name=source_name,
            metadata=metadata,
        )
        attach_source_evidence(
            source,
            location={**location, "row_count": len(rows)},
            text=f"{source_name}: {len(rows)} rows",
        )
        return source


class APIResponseLoader:
    """Load a captured or live JSON API response as source data."""

    name = "api_response_loader"
    source_type = "api_response"

    def __init__(
        self,
        *,
        endpoint_url: str | None = None,
        headers: Mapping[str, str] | None = None,
        timeout_seconds: float = 30.0,
        record_path: str = "$.data",
    ) -> None:
        self.endpoint_url = endpoint_url
        self.headers = dict(headers or {})
        self.timeout_seconds = timeout_seconds
        self.record_path = record_path

    def load(self, path: str | Path) -> list[SourceRecord]:
        source_name = str(self.endpoint_url or path)
        if self.endpoint_url or str(path).startswith(("http://", "https://")):
            url = self.endpoint_url or str(path)
            assert url is not None
            payload = self._fetch(url)
            location = {"kind": "api_response", "url": url}
        else:
            response_path = Path(path)
            payload = _load_json_payload(response_path)
            location = file_location(response_path, kind="api_response_file")

        records = self._records_from_payload(payload)
        metadata: JsonDict = {
            "loader": self.name,
            "record_path": self.record_path if records else None,
            "row_count": len(records),
            "json_payload": payload,
        }
        if records:
            source = SourceRecord.from_rows(
                normalize_rows(records),
                source_type=self.source_type,
                source_name=source_name,
                metadata=metadata,
            )
        else:
            source = SourceRecord.from_text(
                json.dumps(payload, indent=2, sort_keys=True, default=str),
                source_type=self.source_type,
                source_name=source_name,
                metadata=metadata,
            )
        attach_source_evidence(
            source,
            location=location,
            text=f"{source_name}: API response",
            metadata={"record_count": len(records)},
        )
        return [source]

    def _fetch(self, url: str) -> Any:
        request = urllib.request.Request(url, headers=self.headers)
        with urllib.request.urlopen(request, timeout=self.timeout_seconds) as response:
            charset = response.headers.get_content_charset() or "utf-8"
            return json.loads(response.read().decode(charset))

    @staticmethod
    def _records_from_payload(payload: Any) -> list[JsonDict]:
        if isinstance(payload, list) and all(isinstance(item, dict) for item in payload):
            return [dict(item) for item in payload]
        if isinstance(payload, dict) and isinstance(payload.get("data"), list):
            items = payload["data"]
            if all(isinstance(item, dict) for item in items):
                return [dict(item) for item in items]
        return []


class StructuredAPILoader(APIResponseLoader):
    """Alias for API payloads that are already shaped as structured records."""

    name = "structured_api_loader"
    source_type = "structured_api"


class XMLLoader:
    """Load XML documents into flattened structured rows with path metadata."""

    name = "xml_loader"
    source_type = "xml"

    def load(self, path: str | Path) -> list[SourceRecord]:
        xml_path = Path(path)
        tree = ET.parse(xml_path)
        root = tree.getroot()
        record_elements = self._record_elements(root)
        row_paths: list[str] = []
        rows: list[JsonDict] = []
        for index, element in enumerate(record_elements, start=1):
            row_paths.append(self._element_path(root, element, index))
            row = self._flatten_element(element)
            if row:
                rows.append(row)

        if not rows:
            rows = [self._flatten_element(root)]
            row_paths = [f"/{_local_name(root.tag)}"]

        source = SourceRecord.from_rows(
            normalize_rows(rows),
            source_type=self.source_type,
            source_name=str(xml_path),
            metadata={
                "loader": self.name,
                "root_tag": _local_name(root.tag),
                "row_count": len(rows),
                "row_paths": row_paths,
            },
        )
        attach_source_evidence(
            source,
            location=file_location(
                xml_path,
                kind="xml_document",
                root_tag=_local_name(root.tag),
                row_paths=row_paths,
            ),
            text=f"{_local_name(root.tag)}: {len(rows)} XML records",
        )
        return [source]

    @staticmethod
    def _record_elements(root: ET.Element) -> list[ET.Element]:
        children = list(root)
        if not children:
            return [root]
        tag_counts: dict[str, int] = {}
        for child in children:
            tag_counts[_local_name(child.tag)] = tag_counts.get(_local_name(child.tag), 0) + 1
        repeated = {tag for tag, count in tag_counts.items() if count > 1}
        if len(tag_counts) == 1 or repeated:
            return [child for child in children if _local_name(child.tag) in repeated] or children
        return children

    def _flatten_element(
        self,
        element: ET.Element,
        prefix: str = "",
        *,
        name_override: str | None = None,
    ) -> JsonDict:
        row: JsonDict = {}
        name = name_override or _local_name(element.tag)
        key_prefix = f"{prefix}.{name}" if prefix else name

        for attr_name, attr_value in sorted(element.attrib.items()):
            row[f"{key_prefix}.@{_local_name(attr_name)}"] = attr_value

        children = list(element)
        text_value = (element.text or "").strip()
        if text_value and not children:
            row[key_prefix] = text_value

        child_counts: dict[str, int] = {}
        for child in children:
            child_name = _local_name(child.tag)
            child_counts[child_name] = child_counts.get(child_name, 0) + 1

        child_seen: dict[str, int] = {}
        for child in children:
            child_name = _local_name(child.tag)
            child_seen[child_name] = child_seen.get(child_name, 0) + 1
            if child_counts[child_name] > 1:
                child_name = f"{child_name}[{child_seen[child_name] - 1}]"
            row.update(self._flatten_element(child, key_prefix, name_override=child_name))
        return row

    @staticmethod
    def _element_path(root: ET.Element, element: ET.Element, index: int) -> str:
        if element is root:
            return f"/{_local_name(root.tag)}"
        return f"/{_local_name(root.tag)}/{_local_name(element.tag)}[{index}]"


class _HTMLDocumentParser(HTMLParser):
    """Small stdlib HTML parser for text, metadata, links, headings, and JSON-LD."""

    def __init__(self, *, base_url: str | None = None) -> None:
        super().__init__(convert_charrefs=True)
        self.base_url = base_url
        self.title_parts: list[str] = []
        self.visible_parts: list[str] = []
        self.meta_tags: list[JsonDict] = []
        self.links: list[JsonDict] = []
        self.headings: list[JsonDict] = []
        self.json_ld: list[JsonDict] = []
        self._in_title = False
        self._skip_depth = 0
        self._jsonld_depth = 0
        self._jsonld_parts: list[str] = []
        self._current_heading: JsonDict | None = None
        self._current_anchor: JsonDict | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = {key.lower(): value or "" for key, value in attrs}
        tag_name = tag.lower()
        if tag_name in {"script", "style", "noscript"}:
            self._skip_depth += 1
            if tag_name == "script" and "ld+json" in attrs_dict.get("type", "").casefold():
                self._jsonld_depth += 1
                self._jsonld_parts = []
            return
        if tag_name == "title":
            self._in_title = True
        elif tag_name == "meta":
            self._add_meta(attrs_dict)
        elif tag_name == "link":
            self._add_link(attrs_dict)
        elif tag_name == "a":
            self._current_anchor = {
                "href": self._absolute_url(attrs_dict.get("href", "")),
                "text": "",
                "rel": attrs_dict.get("rel", ""),
            }
        elif tag_name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._current_heading = {"level": int(tag_name[1]), "text": ""}

    def handle_endtag(self, tag: str) -> None:
        tag_name = tag.lower()
        if tag_name == "title":
            self._in_title = False
        elif tag_name in {"script", "style", "noscript"}:
            if self._jsonld_depth:
                self._add_jsonld("".join(self._jsonld_parts))
                self._jsonld_depth = 0
                self._jsonld_parts = []
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag_name == "a" and self._current_anchor is not None:
            anchor = {
                **self._current_anchor,
                "text": _collapse_space(str(self._current_anchor.get("text", ""))),
            }
            if anchor["href"] or anchor["text"]:
                self.links.append(anchor)
            self._current_anchor = None
        elif tag_name in {"h1", "h2", "h3", "h4", "h5", "h6"} and self._current_heading:
            heading = {
                **self._current_heading,
                "text": _collapse_space(str(self._current_heading.get("text", ""))),
            }
            if heading["text"]:
                self.headings.append(heading)
            self._current_heading = None

    def handle_data(self, data: str) -> None:
        if self._jsonld_depth:
            self._jsonld_parts.append(data)
            return
        if self._in_title:
            self.title_parts.append(data)
        if self._skip_depth:
            return
        text = _collapse_space(data)
        if not text:
            return
        self.visible_parts.append(text)
        if self._current_anchor is not None:
            self._current_anchor["text"] = f"{self._current_anchor.get('text', '')} {text}".strip()
        if self._current_heading is not None:
            self._current_heading["text"] = f"{self._current_heading.get('text', '')} {text}".strip()

    @property
    def title(self) -> str:
        return _collapse_space(" ".join(self.title_parts))

    @property
    def visible_text(self) -> str:
        return _collapse_space(" ".join(self.visible_parts))

    @property
    def canonical_url(self) -> str:
        for link in self.links:
            rel_values = str(link.get("rel", "")).casefold().split()
            if "canonical" in rel_values and link.get("href"):
                return str(link["href"])
        for tag in self.meta_tags:
            if tag.get("key") in {"og:url", "twitter:url"}:
                return str(tag.get("content", ""))
        return self.base_url or ""

    def metadata_payload(self) -> JsonDict:
        return {
            "title": self.title,
            "description": self.meta_content("description"),
            "canonical_url": self.canonical_url,
            "meta_tags": self.meta_tags,
            "links": self.links,
            "headings": self.headings,
            "json_ld": self.json_ld,
        }

    def meta_content(self, key: str) -> str:
        normalized = key.casefold()
        for tag in self.meta_tags:
            if str(tag.get("key", "")).casefold() == normalized:
                return str(tag.get("content", ""))
        return ""

    def _add_meta(self, attrs: JsonDict) -> None:
        key = attrs.get("name") or attrs.get("property") or attrs.get("itemprop") or attrs.get("http-equiv")
        content = attrs.get("content", "")
        if not key or not content:
            return
        attribute = (
            "name"
            if attrs.get("name")
            else "property"
            if attrs.get("property")
            else "itemprop"
            if attrs.get("itemprop")
            else "http-equiv"
        )
        self.meta_tags.append(
            {
                "key": str(key).strip(),
                "content": _collapse_space(str(content)),
                "attribute": attribute,
                "snippet": f'<meta {attribute}="{key}" content="{content}">',
            }
        )

    def _add_link(self, attrs: JsonDict) -> None:
        href = attrs.get("href", "")
        rel = attrs.get("rel", "")
        if not href and not rel:
            return
        self.links.append(
            {
                "href": self._absolute_url(str(href)),
                "rel": str(rel),
                "type": attrs.get("type", ""),
                "title": attrs.get("title", ""),
            }
        )

    def _add_jsonld(self, raw: str) -> None:
        stripped = raw.strip()
        if not stripped:
            return
        try:
            payload = json.loads(stripped)
        except json.JSONDecodeError:
            self.json_ld.append({"raw": stripped, "parse_error": True})
            return
        items = payload if isinstance(payload, list) else [payload]
        self.json_ld.extend(item for item in items if isinstance(item, dict))

    def _absolute_url(self, value: str) -> str:
        if not value:
            return ""
        if not self.base_url:
            return value
        return urllib.parse.urljoin(self.base_url, value)


class HTMLLoader:
    """Load local HTML documents with visible text and metadata preserved."""

    name = "html_loader"
    source_type = "html"

    def load(self, path: str | Path) -> list[SourceRecord]:
        html_path = Path(path)
        html = html_path.read_text(encoding="utf-8")
        source = self._source_from_html(
            html,
            source_name=str(html_path),
            location=file_location(html_path, kind="html_document"),
            source_type=self.source_type,
            base_url=html_path.resolve().as_uri(),
            extra_metadata={"path": str(html_path), "suffix": html_path.suffix.lower()},
        )
        return [source]

    def _source_from_html(
        self,
        html: str,
        *,
        source_name: str,
        location: JsonDict,
        source_type: str,
        base_url: str | None,
        extra_metadata: JsonDict | None = None,
    ) -> SourceRecord:
        parser = _HTMLDocumentParser(base_url=base_url)
        parser.feed(html)
        html_metadata = parser.metadata_payload()
        content = parser.visible_text or html_metadata.get("title") or source_name
        source = SourceRecord.from_text(
            str(content),
            source_type=source_type,
            source_name=source_name,
            metadata={
                "loader": self.name,
                "title": html_metadata["title"],
                "description": html_metadata["description"],
                "canonical_url": html_metadata["canonical_url"],
                "html_metadata": html_metadata,
                "raw_html_character_count": len(html),
                "visible_character_count": len(str(content)),
                **(extra_metadata or {}),
            },
        )
        attach_source_evidence(
            source,
            location=location,
            text=str(html_metadata["title"] or source_name),
            metadata={
                "meta_tag_count": len(html_metadata["meta_tags"]),
                "link_count": len(html_metadata["links"]),
                "json_ld_count": len(html_metadata["json_ld"]),
            },
        )
        return source


class WebPageLoader(HTMLLoader):
    """Load a web page URL or local HTML fixture as a webpage source."""

    name = "webpage_loader"
    source_type = "webpage"

    def __init__(
        self,
        *,
        headers: Mapping[str, str] | None = None,
        timeout_seconds: float = 30.0,
        user_agent: str = "kg-construction/0.1",
    ) -> None:
        self.headers = {"User-Agent": user_agent, **dict(headers or {})}
        self.timeout_seconds = timeout_seconds

    def load(self, path: str | Path) -> list[SourceRecord]:
        target = str(path)
        if target.startswith(("http://", "https://")):
            html, response_metadata = self._fetch_page(target)
            source = self._source_from_html(
                html,
                source_name=str(response_metadata.get("final_url") or target),
                location={
                    "kind": "webpage",
                    "url": target,
                    "final_url": response_metadata.get("final_url") or target,
                    "status": response_metadata.get("status"),
                },
                source_type=self.source_type,
                base_url=str(response_metadata.get("final_url") or target),
                extra_metadata={"url": target, **response_metadata},
            )
            return [source]
        html_path = Path(path)
        html = html_path.read_text(encoding="utf-8")
        return [
            self._source_from_html(
                html,
                source_name=str(html_path),
                location=file_location(html_path, kind="webpage_file"),
                source_type=self.source_type,
                base_url=html_path.resolve().as_uri(),
                extra_metadata={"path": str(html_path), "suffix": html_path.suffix.lower()},
            )
        ]

    def _fetch_page(self, url: str) -> tuple[str, JsonDict]:
        request = urllib.request.Request(url, headers=self.headers)
        with urllib.request.urlopen(request, timeout=self.timeout_seconds) as response:
            charset = response.headers.get_content_charset() or "utf-8"
            html = response.read().decode(charset, errors="replace")
            metadata = {
                "final_url": response.geturl(),
                "status": getattr(response, "status", None),
                "content_type": response.headers.get("Content-Type", ""),
            }
        return html, metadata


class LogLoader:
    """Load JSON, Apache, key-value, and generic text logs as event rows."""

    name = "log_loader"
    source_type = "logs"

    def load(self, path: str | Path) -> list[SourceRecord]:
        log_path = Path(path)
        rows: list[JsonDict] = []
        parser_counts: dict[str, int] = {}
        for line_number, line in enumerate(log_path.read_text(encoding="utf-8").splitlines(), start=1):
            if not line.strip():
                continue
            row = self._parse_line(line, line_number)
            template = self._mine_template(str(row.get("message") or row.get("raw") or ""))
            row["template"] = template
            row["template_id"] = stable_log_template_id(template)
            parser_name = str(row.get("parser", "generic"))
            parser_counts[parser_name] = parser_counts.get(parser_name, 0) + 1
            rows.append(row)

        source = SourceRecord.from_rows(
            normalize_rows(rows),
            source_type=self.source_type,
            source_name=str(log_path),
            metadata={
                "loader": self.name,
                "row_count": len(rows),
                "parser_counts": parser_counts,
                "template_count": len({row.get("template_id") for row in rows}),
                "line_count": len(log_path.read_text(encoding="utf-8").splitlines()),
            },
        )
        attach_source_evidence(
            source,
            location=file_location(log_path, kind="log_file", row_count=len(rows)),
            text=f"{log_path.name}: {len(rows)} log events",
            metadata={"parser_counts": parser_counts},
        )
        return [source]

    def _parse_line(self, line: str, line_number: int) -> JsonDict:
        stripped = line.strip()
        if stripped.startswith("{") and stripped.endswith("}"):
            try:
                payload = json.loads(stripped)
            except json.JSONDecodeError:
                pass
            else:
                if isinstance(payload, dict):
                    return self._json_log_row(payload, line_number, stripped)

        if match := _APACHE_LOG_RE.match(stripped):
            row = {key: value for key, value in match.groupdict().items() if value not in (None, "")}
            status = int(str(row.get("status", "0")))
            row.update(
                {
                    "line_number": line_number,
                    "raw": stripped,
                    "parser": "apache_combined" if row.get("user_agent") is not None else "apache_common",
                    "level": "ERROR" if status >= 500 else "WARNING" if status >= 400 else "INFO",
                    "message": f"{row.get('method', '')} {row.get('url', '')} -> {status}",
                }
            )
            return row

        key_values = {
            match.group("key"): match.group("value").strip('"')
            for match in _KEY_VALUE_RE.finditer(stripped)
        }
        if key_values:
            return {
                **key_values,
                "line_number": line_number,
                "raw": stripped,
                "parser": "key_value",
                "timestamp": key_values.get("timestamp") or key_values.get("time") or key_values.get("ts"),
                "level": str(key_values.get("level", key_values.get("severity", "INFO"))).upper(),
                "message": key_values.get("message") or key_values.get("msg") or stripped,
            }

        generic = _GENERIC_LOG_RE.match(stripped)
        groups = generic.groupdict() if generic else {}
        return {
            "line_number": line_number,
            "raw": stripped,
            "parser": "generic",
            "timestamp": groups.get("timestamp") if groups else None,
            "level": str(groups.get("level") or "INFO").upper(),
            "logger": groups.get("logger") if groups else None,
            "message": groups.get("message") if groups else stripped,
        }

    @staticmethod
    def _json_log_row(payload: JsonDict, line_number: int, raw: str) -> JsonDict:
        timestamp = payload.get("@timestamp") or payload.get("timestamp") or payload.get("time")
        level = payload.get("level") or payload.get("severity") or payload.get("log.level") or "INFO"
        message = payload.get("message") or payload.get("msg") or payload.get("event.action") or raw
        return {
            **payload,
            "line_number": line_number,
            "raw": raw,
            "parser": "json",
            "timestamp": timestamp,
            "level": str(level).upper(),
            "message": message,
        }

    @staticmethod
    def _mine_template(message: str) -> str:
        template = _TEMPLATE_UUID_RE.sub("<UUID>", message)
        template = _TEMPLATE_HEX_RE.sub("<HEX>", template)
        template = _TEMPLATE_ID_RE.sub("<ID>", template)
        template = _TEMPLATE_NUMBER_RE.sub("<NUM>", template)
        return _collapse_space(template)


def stable_log_template_id(template: str) -> str:
    from kg_construction.domain.models import stable_id

    return stable_id("logtemplate", template, length=12)


class ScientificPaperLoader:
    """Load scientific paper XML/PDF/text into paper metadata and narrative content."""

    name = "scientific_paper_loader"
    source_type = "scientific_paper"

    def load(self, path: str | Path) -> list[SourceRecord]:
        paper_path = Path(path)
        suffix = paper_path.suffix.lower()
        if suffix in {".xml", ".tei", ".jats"}:
            source = self._load_xml_paper(paper_path)
        elif suffix == ".pdf":
            source = self._load_pdf_paper(paper_path)
        else:
            source = self._load_text_paper(paper_path)
        return [source]

    def _load_xml_paper(self, path: Path) -> SourceRecord:
        root = ET.parse(path).getroot()
        metadata = self._paper_metadata_from_xml(root)
        text = self._paper_text(metadata)
        source = SourceRecord.from_text(
            text,
            source_type=self.source_type,
            source_name=str(path),
            metadata={"loader": self.name, "paper_format": metadata.pop("paper_format"), **metadata},
        )
        attach_source_evidence(
            source,
            location=file_location(
                path,
                kind="scientific_paper_xml",
                paper_format=source.metadata.get("paper_format"),
            ),
            text=str(source.metadata.get("title") or path.name),
            metadata={"doi": source.metadata.get("doi"), "section_count": len(source.metadata.get("sections", []))},
        )
        return source

    def _load_pdf_paper(self, path: Path) -> SourceRecord:
        pdf_source = PDFLoader().load(path)[0]
        metadata = self._paper_metadata_from_text(pdf_source.content or "", paper_format="pdf")
        source = SourceRecord.from_text(
            pdf_source.content or "",
            source_type=self.source_type,
            source_name=str(path),
            metadata={
                "loader": self.name,
                **metadata,
                "pdf_metadata": pdf_source.metadata,
            },
        )
        attach_source_evidence(
            source,
            location=file_location(path, kind="scientific_paper_pdf"),
            text=str(metadata.get("title") or path.name),
            metadata={"doi": metadata.get("doi"), "section_count": len(metadata.get("sections", []))},
        )
        return source

    def _load_text_paper(self, path: Path) -> SourceRecord:
        text = path.read_text(encoding="utf-8")
        metadata = self._paper_metadata_from_text(text, paper_format="text")
        source = SourceRecord.from_text(
            text,
            source_type=self.source_type,
            source_name=str(path),
            metadata={"loader": self.name, **metadata},
        )
        attach_source_evidence(
            source,
            location=file_location(path, kind="scientific_paper_text"),
            text=str(metadata.get("title") or path.name),
            metadata={"doi": metadata.get("doi"), "section_count": len(metadata.get("sections", []))},
        )
        return source

    def _paper_metadata_from_xml(self, root: ET.Element) -> JsonDict:
        paper_format = "tei" if _local_name(root.tag).casefold() == "tei" else "jats"
        if paper_format == "tei":
            return self._tei_metadata(root)
        return self._jats_metadata(root)

    def _jats_metadata(self, root: ET.Element) -> JsonDict:
        title = _first_element_text(root, "article-title")
        abstract = _first_element_text(root, "abstract")
        authors = []
        for contrib in _elements(root, "contrib"):
            if contrib.attrib.get("contrib-type", "author") != "author":
                continue
            name = _name_from_element(contrib)
            if name:
                authors.append(name)
        doi = ""
        for article_id in _elements(root, "article-id"):
            if article_id.attrib.get("pub-id-type", "").casefold() == "doi":
                doi = _element_text(article_id)
                break
        sections = self._xml_sections(root)
        return {
            "paper_format": "jats",
            "title": title,
            "abstract": abstract,
            "authors": authors,
            "doi": doi or _first_doi(_element_text(root)),
            "journal": _first_element_text(root, "journal-title"),
            "published": _published_date(root),
            "sections": sections,
            "citations": self._xml_citations(root),
        }

    def _tei_metadata(self, root: ET.Element) -> JsonDict:
        title = _first_element_text(root, "title", attr_name="type", attr_value="main") or _first_element_text(root, "title")
        abstract = _first_element_text(root, "abstract")
        authors = [_name_from_element(author) for author in _elements(root, "author")]
        doi = ""
        for idno in _elements(root, "idno"):
            if idno.attrib.get("type", "").casefold() == "doi":
                doi = _element_text(idno)
                break
        return {
            "paper_format": "tei",
            "title": title,
            "abstract": abstract,
            "authors": [author for author in authors if author],
            "doi": doi or _first_doi(_element_text(root)),
            "journal": _first_element_text(root, "title", attr_name="level", attr_value="j"),
            "published": _first_element_text(root, "date"),
            "sections": self._xml_sections(root),
            "citations": self._xml_citations(root),
        }

    @staticmethod
    def _xml_sections(root: ET.Element) -> list[JsonDict]:
        sections: list[JsonDict] = []
        for index, element in enumerate(_elements(root, "sec"), start=1):
            title = _first_child_text(element, "title") or f"Section {index}"
            text = _element_text(element)
            sections.append({"index": index, "title": title, "text": text})
        if sections:
            return sections
        for index, element in enumerate(_elements(root, "div"), start=1):
            title = _first_child_text(element, "head") or _first_child_text(element, "title") or f"Section {index}"
            sections.append({"index": index, "title": title, "text": _element_text(element)})
        return sections

    @staticmethod
    def _xml_citations(root: ET.Element) -> list[JsonDict]:
        citations: list[JsonDict] = []
        for index, ref in enumerate([*_elements(root, "ref"), *_elements(root, "biblStruct")], start=1):
            text = _element_text(ref)
            if text:
                citations.append({"index": index, "text": text[:500], "doi": _first_doi(text)})
        return citations

    @staticmethod
    def _paper_metadata_from_text(text: str, *, paper_format: str) -> JsonDict:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        title = lines[0] if lines else ""
        abstract = ""
        for index, line in enumerate(lines):
            if line.casefold() == "abstract" and index + 1 < len(lines):
                abstract = lines[index + 1]
                break
            if line.casefold().startswith("abstract"):
                abstract = line.split(":", 1)[-1].strip()
                break
        sections = [
            {"index": index, "title": line, "text": ""}
            for index, line in enumerate(lines, start=1)
            if line.casefold() in {"abstract", "introduction", "methods", "results", "discussion", "conclusion"}
        ]
        return {
            "paper_format": paper_format,
            "title": title,
            "abstract": abstract,
            "authors": [],
            "doi": _first_doi(text),
            "journal": "",
            "published": "",
            "sections": sections,
            "citations": [],
        }

    @staticmethod
    def _paper_text(metadata: JsonDict) -> str:
        fragments = [
            str(metadata.get("title", "")),
            str(metadata.get("abstract", "")),
            *[str(section.get("text", "")) for section in metadata.get("sections", []) if isinstance(section, dict)],
        ]
        return "\n\n".join(fragment for fragment in fragments if fragment.strip())


def _local_name(tag: str) -> str:
    if "}" in tag:
        return tag.rsplit("}", 1)[1]
    return tag.split(":", 1)[-1]


def _elements(root: ET.Element, local_name: str) -> list[ET.Element]:
    return [element for element in root.iter() if _local_name(element.tag) == local_name]


def _element_text(element: ET.Element) -> str:
    return _collapse_space(" ".join(text for text in element.itertext() if text and text.strip()))


def _first_element_text(
    root: ET.Element,
    local_name: str,
    *,
    attr_name: str | None = None,
    attr_value: str | None = None,
) -> str:
    for element in _elements(root, local_name):
        if attr_name and element.attrib.get(attr_name) != attr_value:
            continue
        text = _element_text(element)
        if text:
            return text
    return ""


def _first_child_text(element: ET.Element, local_name: str) -> str:
    for child in list(element):
        if _local_name(child.tag) == local_name:
            return _element_text(child)
    return ""


def _name_from_element(element: ET.Element) -> str:
    surname = _first_element_text(element, "surname")
    given = _first_element_text(element, "given-names") or _first_element_text(element, "forename")
    if surname or given:
        return _collapse_space(f"{given} {surname}")
    return _first_element_text(element, "string-name") or _element_text(element)


def _published_date(root: ET.Element) -> str:
    pub_dates = _elements(root, "pub-date")
    if not pub_dates:
        return ""
    pub_date = pub_dates[0]
    year = _first_child_text(pub_date, "year")
    month = _first_child_text(pub_date, "month")
    day = _first_child_text(pub_date, "day")
    return "-".join(part for part in [year, month.zfill(2) if month else "", day.zfill(2) if day else ""] if part)


def _first_doi(text: str) -> str:
    match = _DOI_RE.search(text)
    return match.group(0).rstrip(".") if match else ""


class MarkdownLoader:
    """Load Markdown documents with heading metadata preserved."""

    name = "markdown_loader"
    source_type = "markdown"

    def load(self, path: str | Path) -> list[SourceRecord]:
        markdown_path = Path(path)
        text = markdown_path.read_text(encoding="utf-8")
        headings = [
            {
                "line_number": line_number,
                "level": len(match.group("level")),
                "title": match.group("title"),
            }
            for line_number, line in enumerate(text.splitlines(), start=1)
            if (match := _MARKDOWN_HEADING_RE.match(line))
        ]
        source = SourceRecord.from_text(
            text,
            source_type=self.source_type,
            source_name=str(markdown_path),
            metadata={
                "loader": self.name,
                "suffix": markdown_path.suffix.lower(),
                "headings": headings,
                "heading_count": len(headings),
            },
        )
        attach_source_evidence(
            source,
            location=file_location(
                markdown_path,
                kind="markdown_document",
                heading_count=len(headings),
            ),
            text=str(headings[0]["title"]) if headings else markdown_path.name,
        )
        return [source]


class ReportLoader:
    """Load plain-text or Markdown reports as narrative source records."""

    name = "report_loader"
    source_type = "report"

    def load(self, path: str | Path) -> list[SourceRecord]:
        report_path = Path(path)
        text = report_path.read_text(encoding="utf-8")
        sections = self._sections(text)
        source = SourceRecord.from_text(
            text,
            source_type=self.source_type,
            source_name=str(report_path),
            metadata={
                "loader": self.name,
                "suffix": report_path.suffix.lower(),
                "sections": sections,
                "section_count": len(sections),
            },
        )
        attach_source_evidence(
            source,
            location=file_location(report_path, kind="report_document", section_count=len(sections)),
            text=sections[0]["title"] if sections else report_path.name,
        )
        return [source]

    @staticmethod
    def _sections(text: str) -> list[JsonDict]:
        sections: list[JsonDict] = []
        for line_number, line in enumerate(text.splitlines(), start=1):
            stripped = line.strip()
            heading = _MARKDOWN_HEADING_RE.match(stripped)
            if heading:
                sections.append(
                    {
                        "line_number": line_number,
                        "level": len(heading.group("level")),
                        "title": heading.group("title"),
                    }
                )
            elif stripped and stripped == stripped.title() and len(stripped.split()) <= 8:
                sections.append({"line_number": line_number, "level": 1, "title": stripped})
        return sections


class PDFLoader:
    """Load digitally-born PDF text with page-level metadata."""

    name = "pdf_loader"
    source_type = "pdf"

    def load(self, path: str | Path) -> list[SourceRecord]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover - exercised when dependency is absent
            raise _optional_dependency("pypdf", "Install kg-construction with PDF support.") from exc

        pdf_path = Path(path)
        reader = PdfReader(str(pdf_path))
        page_texts: list[str] = []
        page_metadata: list[JsonDict] = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            page_texts.append(text)
            page_metadata.append({"page_number": index, "char_count": len(text)})

        text = "\n\n".join(page.strip() for page in page_texts if page.strip())
        if not text:
            text = self._extract_literal_pdf_strings(pdf_path)

        document_metadata = {
            key.lstrip("/"): str(value)
            for key, value in (reader.metadata or {}).items()
            if value is not None
        }
        source = SourceRecord.from_text(
            text,
            source_type=self.source_type,
            source_name=str(pdf_path),
            metadata={
                "loader": self.name,
                "page_count": len(reader.pages),
                "pages": page_metadata,
                "document_metadata": document_metadata,
            },
        )
        attach_source_evidence(
            source,
            location=file_location(pdf_path, kind="pdf_document", page_count=len(reader.pages)),
            text=document_metadata.get("Title") or pdf_path.name,
        )
        return [source]

    @staticmethod
    def _extract_literal_pdf_strings(path: Path) -> str:
        raw = path.read_bytes().decode("latin-1", errors="ignore")
        fragments = [
            match.group("text").replace(r"\(", "(").replace(r"\)", ")").replace(r"\\", "\\")
            for match in _PDF_TEXT_RE.finditer(raw)
        ]
        return "\n".join(fragment for fragment in fragments if fragment.strip())


class EmailLoader:
    """Load RFC 5322 email files and preserve thread/header evidence."""

    name = "email_loader"
    source_type = "email"

    def load(self, path: str | Path) -> list[SourceRecord]:
        email_path = Path(path)
        with email_path.open("rb") as handle:
            message = BytesParser(policy=policy.default).parse(handle)
        headers = self._headers(message)
        body = self._body(message)
        source = SourceRecord.from_text(
            body,
            source_type=self.source_type,
            source_name=str(email_path),
            metadata={
                "loader": self.name,
                "headers": headers,
                "attachment_count": self._attachment_count(message),
            },
        )
        attach_source_evidence(
            source,
            location=file_location(
                email_path,
                kind="email_message",
                message_id=headers.get("message_id"),
                subject=headers.get("subject"),
            ),
            text=headers.get("subject") or email_path.name,
        )
        return [source]

    @staticmethod
    def _headers(message: Message) -> JsonDict:
        return {
            "from": str(message.get("from", "")),
            "to": str(message.get("to", "")),
            "cc": str(message.get("cc", "")),
            "subject": str(message.get("subject", "")),
            "date": str(message.get("date", "")),
            "message_id": str(message.get("message-id", "")),
        }

    @staticmethod
    def _body(message: Message) -> str:
        if message.is_multipart():
            parts: list[str] = []
            for part in message.walk():
                if part.get_content_maintype() == "multipart":
                    continue
                if part.get_content_disposition() == "attachment":
                    continue
                if part.get_content_type() == "text/plain":
                    parts.append(_message_content(part))
            return "\n\n".join(part.strip() for part in parts if part.strip())
        return _message_content(message)

    @staticmethod
    def _attachment_count(message: Message) -> int:
        return sum(
            1
            for part in message.walk()
            if part.get_content_disposition() == "attachment"
        )


def _message_content(message: Message) -> str:
    if isinstance(message, EmailMessage):
        content = message.get_content()
        return content if isinstance(content, str) else str(content)
    payload = message.get_payload(decode=True)
    if isinstance(payload, bytes):
        charset = message.get_content_charset() or "utf-8"
        return payload.decode(charset, errors="replace")
    raw_payload = message.get_payload()
    return raw_payload if isinstance(raw_payload, str) else str(raw_payload)


class ChatTranscriptLoader:
    """Load chat transcripts as turn-level structured records."""

    name = "chat_transcript_loader"
    source_type = "chat_transcript"

    def load(self, path: str | Path) -> list[SourceRecord]:
        transcript_path = Path(path)
        text = transcript_path.read_text(encoding="utf-8")
        messages = self._messages(text)
        source = SourceRecord.from_rows(
            normalize_rows(messages),
            source_type=self.source_type,
            source_name=str(transcript_path),
            metadata={
                "loader": self.name,
                "message_count": len(messages),
                "speakers": sorted(
                    {str(row["speaker"]) for row in messages if row.get("speaker")}
                ),
            },
        )
        attach_source_evidence(
            source,
            location=file_location(
                transcript_path,
                kind="chat_transcript",
                message_count=len(messages),
            ),
            text=f"{transcript_path.name}: {len(messages)} messages",
        )
        return [source]

    @staticmethod
    def _messages(text: str) -> list[JsonDict]:
        messages: list[JsonDict] = []
        for line_number, line in enumerate(text.splitlines(), start=1):
            stripped = line.strip()
            if not stripped:
                continue
            timestamp_match = _CHAT_WITH_TIMESTAMP_RE.match(stripped)
            speaker_match = _CHAT_SPEAKER_RE.match(stripped)
            if timestamp_match:
                messages.append(
                    {
                        "line_number": line_number,
                        "timestamp": timestamp_match.group("timestamp"),
                        "speaker": timestamp_match.group("speaker").strip(),
                        "message": timestamp_match.group("message").strip(),
                    }
                )
            elif speaker_match:
                messages.append(
                    {
                        "line_number": line_number,
                        "timestamp": None,
                        "speaker": speaker_match.group("speaker").strip(),
                        "message": speaker_match.group("message").strip(),
                    }
                )
            else:
                messages.append(
                    {
                        "line_number": line_number,
                        "timestamp": None,
                        "speaker": "unknown",
                        "message": stripped,
                    }
                )
        return messages


__all__ = [
    "APIResponseLoader",
    "ChatTranscriptLoader",
    "EmailLoader",
    "ExcelLoader",
    "HTMLLoader",
    "LogLoader",
    "MarkdownLoader",
    "PDFLoader",
    "ReportLoader",
    "SQLDatabaseLoader",
    "StructuredAPILoader",
    "ScientificPaperLoader",
    "WebPageLoader",
    "XMLLoader",
]
